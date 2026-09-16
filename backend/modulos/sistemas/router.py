import io
import os
from datetime import datetime
import pandas as pd
from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile, File
from fastapi.responses import StreamingResponse, Response
from sqlalchemy.orm import Session

# Importamos las dependencias desde tu módulo patrimonio (o global) que ya funciona
from backend.modulos.patrimonio.dependencies import get_db, get_current_user
JWT_DEPENDENCY = get_current_user

from . import schemas
from .models import NotaSalida
from .service import SistemasService
from backend.modulos.auditoria.service import registrar_evento

router = APIRouter(
    prefix="/sistemas",
    tags=["Sistemas"]
)

DIR_NOTAS = os.path.join("frontend", "notas")

# 2. El endpoint debe usar @router.post, NO @app.post
@router.post("/nota-salida/generar")
def generar_nota_salida_pdf(
    req: schemas.NotaSalidaRequest,
    request: Request,
    db: Session = Depends(get_db),
    usuario_actual: str = Depends(JWT_DEPENDENCY)
):
    service = SistemasService(db)
    pdf_buffer, resuelto = service.generar_pdf_salida(req)

    from backend.modulos.usuarios.models import User
    user = db.query(User).filter(User.email == usuario_actual).first()

    # Número secuencial de la nota
    cantidad = db.query(NotaSalida).count() + 1
    numero = f"NS-{cantidad:04d}"

    # Guardamos el PDF en disco para permitir reimpresión
    os.makedirs(DIR_NOTAS, exist_ok=True)
    ruta_pdf = os.path.join(DIR_NOTAS, f"{numero}.pdf")
    with open(ruta_pdf, "wb") as f:
        f.write(pdf_buffer.getvalue())

    nota = NotaSalida(
        numero=numero,
        creador_email=user.email if user else usuario_actual,
        creador_nombre=f"{user.nombre} {user.apellido}".strip() if user else None,
        nombre_apellido=req.nombre_apellido,
        cuil_cuit=req.cuil_cuit,
        motivo=req.motivo,
        numero_inventario=req.numero_inventario,
        cantidad=req.cantidad,
        tipo=resuelto["tipo"],
        marca=resuelto["marca"],
        serie=resuelto["serie"],
        archivo_url=f"/archivos/notas/{numero}.pdf",
    )
    db.add(nota)
    db.commit()
    db.refresh(nota)

    if user:
        registrar_evento(
            db, usuario_id=user.id, accion="GENERAR_NOTA_SALIDA",
            detalle=f"Nota {numero} de salida generada para: {req.nombre_apellido} ({req.cuil_cuit})",
            ip_address=request.client.host,
            pc_nombre=request.headers.get("X-PC-Nombre", "Desconocido"),
            pc_usuario=request.headers.get("X-PC-Usuario", "Desconocido")
        )

    nombre_limpio = req.nombre_apellido.replace(' ', '_')
    return StreamingResponse(
        iter([pdf_buffer.getvalue()]),
        media_type="application/pdf",
        headers={"Content-Disposition": f"inline; filename=Nota_Salida_{numero}_{nombre_limpio}.pdf"}
    )

# --- LISTAR HISTORIAL DE NOTAS DE SALIDA (reimpresión / consulta) ---
@router.get("/notas-salida", response_model=list[schemas.NotaSalidaResponse])
def listar_notas_salida(
    usuario_actual: str = Depends(JWT_DEPENDENCY),
    db: Session = Depends(get_db)
):
    return db.query(NotaSalida).order_by(NotaSalida.fecha_creacion.desc(), NotaSalida.id.desc()).all()

# --- IMPRIMIR/REIMPRIMIR PDF DE UNA NOTA GUARDADA ---
@router.get("/nota-salida/{nota_id}/pdf")
def imprimir_nota_pdf(
    nota_id: int,
    usuario_actual: str = Depends(JWT_DEPENDENCY),
    db: Session = Depends(get_db)
):
    nota = db.query(NotaSalida).filter(NotaSalida.id == nota_id).first()
    if not nota:
        raise HTTPException(status_code=404, detail="Nota no encontrada")

    nombre_archivo = f"{nota.numero}.pdf"
    ruta = os.path.join(DIR_NOTAS, nombre_archivo)
    if not os.path.exists(ruta):
        raise HTTPException(status_code=404, detail="El PDF de la nota ya no existe en el servidor")

    with open(ruta, "rb") as f:
        contenido = f.read()

    return StreamingResponse(
        iter([contenido]),
        media_type="application/pdf",
        headers={"Content-Disposition": f"inline; filename={nombre_archivo}"}
    )

# --- EXPORTAR NOTAS DE SALIDA (Excel) ---
@router.get("/notas-salida/exportar")
def exportar_notas_salida(
    usuario_actual: str = Depends(JWT_DEPENDENCY),
    db: Session = Depends(get_db)
):
    notas = db.query(NotaSalida).order_by(NotaSalida.fecha_creacion.desc(), NotaSalida.id.desc()).all()

    data = []
    for n in notas:
        data.append({
            "Nº": n.numero,
            "Fecha": n.fecha_creacion.strftime("%d/%m/%Y %H:%M") if n.fecha_creacion else "",
            "Receptor": n.nombre_apellido,
            "CUIL/CUIT": n.cuil_cuit,
            "Motivo": n.motivo,
            "Nº Inventario": n.numero_inventario or "",
            "Cantidad": n.cantidad or 1,
            "Tipo": n.tipo or "",
            "Marca": n.marca or "",
            "Serie": n.serie or "",
            "Generada por": n.creador_nombre or n.creador_email or "",
            "Archivo": n.archivo_url or "",
        })

    df = pd.DataFrame(data)
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        total_df = pd.DataFrame([{"Total": f"{len(data)} registros"}])
        total_df.to_excel(writer, index=False, sheet_name="Notas de Salida", startrow=0)
        if not df.empty:
            df.to_excel(writer, index=False, sheet_name="Notas de Salida", startrow=2)

        worksheet = writer.sheets["Notas de Salida"]
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        total_cell = worksheet["A1"]
        total_cell.font = Font(bold=True, size=12, color="1a3644")
        total_cell.fill = PatternFill(start_color="E0E6E8", end_color="E0E6E8", fill_type="solid")
        worksheet.merge_cells("A1:B1")

        for col_idx in range(1, worksheet.max_column + 1):
            max_length = 0
            col_letter = get_column_letter(col_idx)
            for row in range(1, worksheet.max_row + 1):
                cell = worksheet.cell(row=row, column=col_idx)
                try:
                    if cell.value is not None:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            if max_length > 0:
                worksheet.column_dimensions[col_letter].width = min(max_length + 2, 45)

    return Response(
        buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Notas_de_Salida.xlsx"}
    )

# --- IMPORTAR NOTAS DE SALIDA (Excel/CSV) ---
def _columna(fila, *nombres):
    for k in nombres:
        v = fila.get(k)
        if v is None:
            continue
        if isinstance(v, float) and pd.isna(v):
            continue
        if str(v).strip() != "":
            return v
    return None

@router.post("/notas-salida/importar")
async def importar_notas_salida(
    file: UploadFile = File(...),
    request: Request = None,
    db: Session = Depends(get_db),
    usuario_actual: str = Depends(JWT_DEPENDENCY)
):
    contenido = await file.read()
    nombre_archivo = (file.filename or "").lower()

    try:
        if nombre_archivo.endswith(".csv"):
            df = pd.read_csv(io.BytesIO(contenido), sep=None, engine="python")
        else:
            df = pd.read_excel(io.BytesIO(contenido))
    except Exception:
        raise HTTPException(status_code=400, detail="No se pudo leer el archivo. Usá .xlsx o .csv con una fila de encabezados.")

    df.columns = [str(c).strip() for c in df.columns]

    from backend.modulos.usuarios.models import User
    user = db.query(User).filter(User.email == usuario_actual).first()
    credor_nombre = f"{user.nombre} {user.apellido}".strip() if user else None

    numeros_existentes = {n.numero for n in db.query(NotaSalida.numero).all()}
    siguiente = db.query(NotaSalida).count() + 1

    importados = 0
    errores = []

    for idx, fila in df.iterrows():
        nombre_apellido = str(_columna(fila, "nombre_apellido", "nombre", "Receptor", "NOMBRE_APELLIDO") or "").strip()
        cuil = str(_columna(fila, "cuil_cuit", "cuil", "cuit", "CUIL/CUIT", "CUIL") or "").strip()
        motivo = str(_columna(fila, "motivo", "Motivo") or "").strip()

        if not nombre_apellido or not cuil or not motivo:
            errores.append(f"Fila {idx + 2}: faltan datos obligatorios (Receptor, CUIL/CUIT, Motivo)")
            continue

        numero = str(_columna(fila, "numero", "Nº", "NRO") or "").strip()
        if not numero:
            while f"NS-{siguiente:04d}" in numeros_existentes:
                siguiente += 1
            numero = f"NS-{siguiente:04d}"
            siguiente += 1
        if numero in numeros_existentes:
            errores.append(f"Fila {idx + 2}: el número {numero} ya existe")
            continue
        numeros_existentes.add(numero)

        fecha = _columna(fila, "fecha_creacion", "fecha", "Fecha")
        if fecha:
            try:
                fecha = pd.Timestamp(fecha).to_pydatetime()
            except Exception:
                fecha = datetime.utcnow()
        else:
            fecha = datetime.utcnow()

        try:
            cantidad = int(_columna(fila, "cantidad", "Cantidad") or 1)
        except Exception:
            cantidad = 1

        nota = NotaSalida(
            numero=numero,
            fecha_creacion=fecha,
            creador_email=user.email if user else usuario_actual,
            creador_nombre=credor_nombre,
            nombre_apellido=nombre_apellido,
            cuil_cuit=cuil,
            motivo=motivo,
            numero_inventario=str(_columna(fila, "numero_inventario", "nro_inventario", "inventario", "Nº Inventario") or "").strip() or None,
            cantidad=cantidad,
            tipo=str(_columna(fila, "tipo", "Tipo") or "").strip() or None,
            marca=str(_columna(fila, "marca", "Marca") or "").strip() or None,
            serie=str(_columna(fila, "serie", "Serie") or "").strip() or None,
            archivo_url=None
        )
        db.add(nota)
        importados += 1

    db.commit()

    if user and importados:
        registrar_evento(
            db, usuario_id=user.id, accion="IMPORTAR_NOTAS_SALIDA",
            detalle=f"Se importaron {importados} notas de salida desde el archivo {file.filename}",
            ip_address=request.client.host if request and request.client else None,
            pc_nombre=request.headers.get("X-PC-Nombre", "Desconocido") if request else "Desconocido",
            pc_usuario=request.headers.get("X-PC-Usuario", "Desconocido") if request else "Desconocido"
        )

    return {"importados": importados, "errores": errores}