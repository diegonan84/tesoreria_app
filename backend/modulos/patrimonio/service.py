from sqlalchemy.orm import Session
from sqlalchemy import cast, Integer, asc, desc, or_, and_
from fastapi import HTTPException, status, UploadFile
from datetime import datetime, timezone
from .models import Patrimonio, PatrimonioHistorial
from .schemas import PatrimonioCreate, PatrimonioUpdate
from .repository import PatrimonioRepository, PatrimonioHistorialRepository
from . import models, schemas
import pandas as pd
import io
import re
import barcode
from barcode.writer import ImageWriter
import base64
from fastapi import Request
from backend.modulos.auth.dependencies import verificar_usuario_autenticado


class PatrimonioService:
    def __init__(self, db: Session):
        self.db = db 
        self.repo = PatrimonioRepository(db)
        self.historial_repo = PatrimonioHistorialRepository(db)

    def _filtros_busqueda_avanzada(self, query, busqueda: str):
        """Búsqueda unificada por Nº de inventario:
        - '100-150'  → rango
        - '100,105,110' o '100;105;110' o mezcla '100,130-140' → lista de números/rangos
        - cualquier otro texto → contiene (como antes)
        """
        if not busqueda:
            return query
        tokens = [t.strip() for t in re.split(r"[;,]+", busqueda) if t.strip()]
        if not tokens:
            return query
        col = cast(models.Patrimonio.numero_inventario, Integer)

        if len(tokens) == 1:
            m = re.fullmatch(r"(\d+)\s*-\s*(\d+)", tokens[0])
            if m:
                a, b = int(m.group(1)), int(m.group(2))
                return query.filter(col >= min(a, b), col <= max(a, b))
            return query.filter(models.Patrimonio.numero_inventario.ilike(f"%{tokens[0]}%"))

        condiciones = []
        for token in tokens:
            m = re.fullmatch(r"(\d+)\s*-\s*(\d+)", token)
            if m:
                a, b = int(m.group(1)), int(m.group(2))
                condiciones.append(and_(col >= min(a, b), col <= max(a, b)))
            elif token.isdigit():
                condiciones.append(col == int(token))
            else:
                condiciones.append(models.Patrimonio.numero_inventario.ilike(f"%{token}%"))
        return query.filter(or_(*condiciones))

    def get_all(self, skip: int = 0, limit: int = 100, estado: str = None, busqueda: str = None, desde: str = None, hasta: str = None, tipo: str = None, anio: str = None, rubro: str = None, sort_by: str = "numero_inventario", orden: str = "asc"):
        query = self.db.query(models.Patrimonio)
        
        # Filtros (Estado, búsqueda y rangos)
        if estado:
            query = query.filter(models.Patrimonio.estado == estado)
        query = self._filtros_busqueda_avanzada(query, busqueda)
        if desde and desde.isdigit():
            query = query.filter(cast(models.Patrimonio.numero_inventario, Integer) >= int(desde))
        if hasta and hasta.isdigit():
            query = query.filter(cast(models.Patrimonio.numero_inventario, Integer) <= int(hasta))
            
        if anio:
            query = query.filter(models.Patrimonio.anio.ilike(f"%{anio}%"))
            
        if rubro:
            query = query.filter(models.Patrimonio.rubro_patrimonial_numero == rubro)

        if tipo:
            if tipo == 'OTROS_GENERAL':
                principales = ['pc', 'notebook', 'scanner', 'impresora']
                query = query.filter(models.Patrimonio.tipo.isnot(None), models.Patrimonio.tipo != '')
                for p in principales:
                    query = query.filter(~models.Patrimonio.tipo.ilike(p))
            else:
                query = query.filter(models.Patrimonio.tipo.ilike(tipo))

        if sort_by == "rubro_patrimonial_numero":
            columna = models.Patrimonio.rubro_patrimonial_numero
        elif sort_by == "rubro_patrimonial_descripcion":
            columna = models.Patrimonio.rubro_patrimonial_descripcion
        elif sort_by == "cuenta":
            columna = models.Patrimonio.cuenta
        elif sort_by == "descripcion_bien":
            columna = models.Patrimonio.descripcion_bien
        elif sort_by == "anio":
            columna = models.Patrimonio.anio
        else:
            columna = cast(models.Patrimonio.numero_inventario, Integer)
            
        if orden == "desc":
            query = query.order_by(desc(columna))
        else:
            query = query.order_by(asc(columna))
            
        return query.offset(skip).limit(limit).all()

    def get_total(self, estado: str = None, busqueda: str = None, desde: str = None, hasta: str = None, tipo: str = None, anio: str = None, rubro: str = None):
        query = self.db.query(models.Patrimonio)
        
        if estado:
            query = query.filter(models.Patrimonio.estado == estado)
        query = self._filtros_busqueda_avanzada(query, busqueda)
            
        if desde and desde.isdigit():
            query = query.filter(cast(models.Patrimonio.numero_inventario, Integer) >= int(desde))
        if hasta and hasta.isdigit():
            query = query.filter(cast(models.Patrimonio.numero_inventario, Integer) <= int(hasta))
            
        if anio:
            query = query.filter(models.Patrimonio.anio.ilike(f"%{anio}%"))
            
        if rubro:
            query = query.filter(models.Patrimonio.rubro_patrimonial_numero == rubro)
            
        if tipo:
            if tipo == 'OTROS_GENERAL':
                principales = ['pc', 'notebook', 'scanner', 'impresora']
                query = query.filter(
                    models.Patrimonio.tipo.isnot(None),
                    models.Patrimonio.tipo != ''
                )
                for p in principales:
                    query = query.filter(~models.Patrimonio.tipo.ilike(p))
            else:
                query = query.filter(models.Patrimonio.tipo.ilike(tipo))
                
        return query.count()

    def get_by_numero(self, numero_inventario: str):
        patrimonio = self.repo.get_by_numero_inventario(numero_inventario)
        if not patrimonio:
            raise HTTPException(status_code=404, detail="Bien patrimonial no encontrado")
        return patrimonio

    def create(self, data: schemas.PatrimonioCreate, usuario: str, ip: str):
        try:
            base_num = int(data.numero_inventario)
        except ValueError:
            raise HTTPException(status_code=400, detail="El número de inventario base debe ser numérico para generar lotes consecutivos.")

        nuevos_bienes = []
        historiales = []

        for i in range(data.cantidad):
            num_actual = str(base_num + i)
            
            existe = self.db.query(models.Patrimonio).filter(models.Patrimonio.numero_inventario == num_actual).first()
            if existe:
                raise HTTPException(status_code=400, detail=f"Conflicto: El número de inventario {num_actual} ya existe. Operación de lote cancelada para proteger los datos.")

            bien_dict = data.model_dump(exclude={"cantidad"})
            bien_dict["numero_inventario"] = num_actual
            bien_dict["usuario_modificacion"] = usuario
            
            nuevo_bien = models.Patrimonio(**bien_dict)
            nuevos_bienes.append(nuevo_bien)

            obs_historial = f"Alta en lote ({data.cantidad} ítems) a partir del número base: {base_num}" if data.cantidad > 1 else "Alta inicial del bien"
            
            nuevo_historial = models.PatrimonioHistorial(
                numero_inventario=num_actual,
                usuario=usuario,
                accion="ALTA",
                ip=ip,
                observaciones=obs_historial
            )
            historiales.append(nuevo_historial)

        try:
            self.db.add_all(nuevos_bienes)
            self.db.add_all(historiales) 
            self.db.commit()
            
            for bien in nuevos_bienes:
                self.db.refresh(bien)
                
            return nuevos_bienes 
            
        except Exception as e:
            self.db.rollback()
            raise HTTPException(status_code=500, detail=f"Error crítico en BD al guardar el lote: {str(e)}")

    def update(self, numero_inventario: str, data: PatrimonioUpdate, usuario: str, ip: str):
        patrimonio = self.get_by_numero(numero_inventario)
        update_data = data.model_dump(exclude_unset=True)
        
        cambios_realizados = False

        for campo, valor_nuevo in update_data.items():
            valor_anterior = getattr(patrimonio, campo)
            
            if valor_anterior != valor_nuevo:
                self._registrar_historial(
                    numero_inventario=patrimonio.numero_inventario,
                    usuario=usuario,
                    accion="MODIFICACION",
                    campo_modificado=campo,
                    valor_anterior=str(valor_anterior) if valor_anterior is not None else "",
                    valor_nuevo=str(valor_nuevo) if valor_nuevo is not None else "",
                    ip=ip
                )
                setattr(patrimonio, campo, valor_nuevo)
                cambios_realizados = True

        if cambios_realizados:
            patrimonio.usuario_modificacion = usuario
            self.repo.update(patrimonio)

        return patrimonio

    def delete(self, numero_inventario: str, usuario: str, ip: str):
        patrimonio = self.get_by_numero(numero_inventario)
        
        if not patrimonio.activo:
            raise HTTPException(status_code=400, detail="El bien ya se encuentra dado de baja")

        self.repo.delete_logico(patrimonio)

        self._registrar_historial(
            numero_inventario=patrimonio.numero_inventario,
            usuario=usuario,
            accion="BAJA",
            campo_modificado="activo",
            valor_anterior="True",
            valor_nuevo="False",
            ip=ip,
            observaciones="Baja lógica del sistema"
        )
        return {"mensaje": "Bien patrimonial dado de baja exitosamente"}

    def get_historial(self, numero_inventario: str):
        self.get_by_numero(numero_inventario)
        return self.historial_repo.get_by_inventario(numero_inventario)

    def _registrar_historial(self, numero_inventario, usuario, accion, ip, campo_modificado=None, valor_anterior=None, valor_nuevo=None, observaciones=None):
        historial = PatrimonioHistorial(
            numero_inventario=numero_inventario,
            usuario=usuario,
            fecha=datetime.now(timezone.utc),
            accion=accion,
            campo_modificado=campo_modificado,
            valor_anterior=valor_anterior,
            valor_nuevo=valor_nuevo,
            ip=ip,
            observaciones=observaciones
        )
        self.historial_repo.create(historial)
    
    async def importar_excel(self, file: UploadFile, usuario: str, ip: str):
        contenido = await file.read()
        try:
            # 1. Cargamos el excel completo sin asumir que la fila 0 son los títulos
            df_completo = pd.read_excel(io.BytesIO(contenido), header=None)
            
            # 2. Buscamos dinámicamente la fila donde están los títulos reales
            header_idx = 0
            for idx, row in df_completo.iterrows():
                row_str = " ".join(str(val).lower() for val in row.values if pd.notna(val))
                # Buscamos palabras clave que siempre están en la fila de títulos
                if "inventario" in row_str and "estado" in row_str:
                    header_idx = idx
                    break
                    
            # 3. Asignamos los títulos correctos y descartamos la "basura" de arriba
            df_completo.columns = df_completo.iloc[header_idx]
            df = df_completo[header_idx + 1:].reset_index(drop=True)
            
            # 4. Normalizamos los nombres de las columnas para evitar problemas de espacios
            df.columns = df.columns.astype(str).str.strip()
            
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"No se pudo leer el Excel. Verifica que no esté corrupto. Error: {str(e)}")

        df = df.where(pd.notnull(df), None)

        altas = 0
        modificaciones = 0
        omitidos = 0
        transferidos = 0

        total_registros = len(df)
        bienes_a_crear = []
        historiales = []
        
        numeros_en_excel = set()

        # Movemos las funciones auxiliares fuera del bucle para mayor velocidad
        def get_val(fila_actual, nombres_columna):
            for col in nombres_columna:
                if col in df.columns:
                    val = fila_actual[col]
                    if pd.isna(val):
                        return None
                    val_str = str(val).strip()
                    if val_str == '' or val_str.lower() == 'nan':
                        return None
                    return val_str
            return None

        def procesar_monto(fila_actual, nombres_columna):
            val = get_val(fila_actual, nombres_columna)
            if val:
                try:
                    val_limpio = val.replace('$', '').replace(',', '.').strip()
                    return float(val_limpio)
                except ValueError:
                    return None
            return None

        for index, row in df.iterrows():
            # Agregamos "Nº Inventario" con el símbolo exacto del Excel
            nro_inv = get_val(row, ['Nº Inventario', 'N  Inventario', 'Nro Inventario', 'numero_inventario', 'Inventario'])
            
            if not nro_inv:
                omitidos += 1
                continue
                
            numeros_en_excel.add(str(nro_inv).strip())

            datos_fila = {
                "institucional": get_val(row, ['Institucional', 'institucional']),
                "cuenta": get_val(row, ['Cuenta', 'cuenta']),
                "estado": get_val(row, ['Estado', 'estado']),
                "numero_migrado": get_val(row, ['Nº Migrado', 'N  Migrado', 'Nro Migrado', 'numero_migrado']),
                "descripcion_item": get_val(row, ['Descripcion del Item', 'Descripción del Item', 'descripcion_item']),
                "descripcion_bien": get_val(row, ['Descripcion del Bien', 'Descripción del Bien', 'descripcion_bien']),
                "marca": get_val(row, ['Marca', 'marca']),
                "modelo": get_val(row, ['Modelo', 'modelo']),
                "serie": get_val(row, ['Serie', 'serie']),
                "anio": get_val(row, ['Año', 'año', 'Anio', 'anio']),
                "reparticion": get_val(row, ['Repartición Usuario', 'Repartición', 'Reparticion', 'reparticion']),
                "usuario": get_val(row, ['Usuario', 'Responsable', 'usuario', 'Repartición Usuario']),
                "observaciones": get_val(row, ['Observaciones', 'observaciones']),
                "monto_original": procesar_monto(row, ['Monto Original', 'monto_original']),
                "monto_actualizado": procesar_monto(row, ['Monto Actualizado', 'monto_actualizado']),
                "monto_residual": procesar_monto(row, ['Monto Residual', 'monto_residual'])
            }

            bien_existente = self.db.query(models.Patrimonio).filter(models.Patrimonio.numero_inventario == nro_inv).first()

            if bien_existente:
                cambios = False
                
                # REGLA ESTRICTA: Solo actualizar estos 4 campos si el bien ya existe
                campos_permitidos_actualizar = ["estado", "monto_original", "monto_actualizado", "monto_residual"]
                
                for campo in campos_permitidos_actualizar:
                    valor_nuevo = datos_fila.get(campo)
                    if valor_nuevo is not None and getattr(bien_existente, campo) != valor_nuevo:
                        setattr(bien_existente, campo, valor_nuevo)
                        cambios = True
                
                if cambios:
                    bien_existente.usuario_modificacion = usuario
                    modificaciones += 1
                    historiales.append(models.PatrimonioHistorial(
                        numero_inventario=nro_inv, usuario=usuario, accion="IMPORTACION_MODIFICACION",
                        ip=ip, observaciones=f"Actualización parcial de montos/estado vía Excel: {file.filename}"
                    ))
                else:
                    omitidos += 1

            else:
                # Si el bien NO existe, se crea con TODOS los campos
                nuevo_bien = models.Patrimonio(
                    numero_inventario=nro_inv,
                    usuario_modificacion=usuario,
                    **datos_fila
                )
                bienes_a_crear.append(nuevo_bien)
                altas += 1
                
                historiales.append(models.PatrimonioHistorial(
                    numero_inventario=nro_inv, usuario=usuario, accion="IMPORTACION_ALTA",
                    ip=ip, observaciones=f"Alta vía Excel: {file.filename}"
                ))

        # LÓGICA DE TRANSFERENCIAS AUTOMÁTICAS POR OMISIÓN EN EXCEL
        bienes_autorizados = self.db.query(models.Patrimonio).filter(models.Patrimonio.estado == 'Autorizado').all()
        
        for bien in bienes_autorizados:
            if bien.numero_inventario not in numeros_en_excel:
                valor_anterior = bien.estado
                bien.estado = 'Transferido'
                bien.usuario_modificacion = usuario
                
                historiales.append(models.PatrimonioHistorial(
                    numero_inventario=bien.numero_inventario,
                    usuario=usuario,
                    accion="MODIFICACION_AUTOMATICA",
                    campo_modificado="estado",
                    valor_anterior=valor_anterior,
                    valor_nuevo="Transferido",
                    ip=ip,
                    observaciones=f"Pasó a Transferido por no estar presente en la nueva importación: {file.filename}"
                ))
                transferidos += 1

        try:
            if bienes_a_crear:
                self.db.add_all(bienes_a_crear)
            if historiales:
                self.db.add_all(historiales)
                
            registro_importacion = models.PatrimonioImportacion(
                usuario=usuario,
                archivo=file.filename,
                cantidad_registros=total_registros,
                cantidad_altas=altas,
                cantidad_modificaciones=modificaciones,
                cantidad_omitidos=omitidos
            )
            self.db.add(registro_importacion)
            
            self.db.commit()

            return {
                "mensaje": "Importación finalizada con éxito",
                "resumen": {
                    "total_filas_leidas": total_registros,
                    "nuevos_creados": altas,
                    "bienes_actualizados": modificaciones,
                    "bienes_transferidos": transferidos,
                    "omitidos_o_sin_cambios": omitidos,
                    "archivo": file.filename
                }
            }
        except Exception as e:
            self.db.rollback()
            raise HTTPException(status_code=500, detail=f"Error crítico guardando en BD. Transacción revertida. Detalle: {str(e)}")

    def exportar_excel(self, busqueda=None, desde=None, hasta=None, anio=None, rubro=None):
        query = self.db.query(models.Patrimonio).filter(models.Patrimonio.estado == "Autorizado")
        
        query = self._filtros_busqueda_avanzada(query, busqueda)
        if desde and desde.isdigit():
            query = query.filter(cast(models.Patrimonio.numero_inventario, Integer) >= int(desde))
        if hasta and hasta.isdigit():
            query = query.filter(cast(models.Patrimonio.numero_inventario, Integer) <= int(hasta))
        if anio:
            query = query.filter(models.Patrimonio.anio.ilike(f"%{anio}%"))
        if rubro:
            query = query.filter(models.Patrimonio.rubro_patrimonial_numero == rubro)
        
        query = query.order_by(cast(models.Patrimonio.numero_inventario, Integer))
        bienes = query.all()
        
        data = []
        for bien in bienes:
            data.append({
                "Nº Inventario": bien.numero_inventario,
                "Descripción del Item": bien.descripcion_item,
                "Descripción del Bien": bien.descripcion_bien,
                "Año": bien.anio,
                "Rubro Patrimonial Número": bien.rubro_patrimonial_numero,
                "Rubro Patrimonial Descripción": bien.rubro_patrimonial_descripcion,
                "Marca": bien.marca,
                "Modelo": bien.modelo,
                "Serie": bien.serie,
                "Responsable": bien.usuario,
                "Repartición": bien.reparticion,
                "Estado": bien.estado,
                "Activo": "Sí" if bien.activo else "Baja",
                "Institucional": bien.institucional,
                "Cuenta": bien.cuenta,
                "Nº Migrado": bien.numero_migrado,
                "Monto Original": float(bien.monto_original) if bien.monto_original else None,
                "Monto Actualizado": float(bien.monto_actualizado) if bien.monto_actualizado else None,
                "Monto Residual": float(bien.monto_residual) if bien.monto_residual else None,
                "Observaciones": bien.observaciones,
            })
            
        df = pd.DataFrame(data)
        
        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
            # Fila de total
            total_df = pd.DataFrame([{"Total": f"{len(data)} registros"}])
            total_df.to_excel(writer, index=False, sheet_name='Inventario', startrow=0)
            
            # Datos a partir de la fila 2
            if not df.empty:
                df.to_excel(writer, index=False, sheet_name='Inventario', startrow=2)
            
            worksheet = writer.sheets['Inventario']
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
            total_cell = worksheet['A1']
            total_cell.font = Font(bold=True, size=12, color="1a3644")
            total_cell.fill = PatternFill(start_color="E0E6E8", end_color="E0E6E8", fill_type="solid")
            worksheet.merge_cells('A1:B1')
            
            for col_idx in range(1, worksheet.max_column + 1):
                max_length = 0
                col_letter = get_column_letter(col_idx)
                for row in range(1, worksheet.max_row + 1):
                    cell = worksheet.cell(row=row, column=col_idx)
                    try:
                        if cell.value is not None:
                            max_length = max(max_length, len(str(cell.value)))
                    except Exception:
                        pass
                if max_length > 0:
                    worksheet.column_dimensions[col_letter].width = min(max_length + 2, 40)

        buffer.seek(0)
        return buffer

    
    def generar_etiquetas(self, numeros_inventario: list[str]):
        bienes = self.db.query(models.Patrimonio).filter(models.Patrimonio.numero_inventario.in_(numeros_inventario)).all()
        
        etiquetas = []
        
        # 1. Cambiamos la clase de código de barras al estándar EAN-13
        CODE_CLASS = barcode.get_barcode_class('ean13')
        
        for bien in bienes:
            # 2. Limpiamos cualquier letra accidental y rellenamos con ceros a la izquierda
            numero_limpio = ''.join(filter(str.isdigit, str(bien.numero_inventario)))
            
            # EAN-13 exige exactamente 12 dígitos de entrada. 
            # zfill(12) agrega los ceros, y [-12:] asegura que no nos pasemos del límite.
            numero_base = numero_limpio.zfill(12)[-12:]
            
            opciones = {
                'module_width': 0.25,
                'module_height': 12.0,
                'font_size': 9,
                'text_distance': 4.0,
                'quiet_zone': 1.0,
                'center_text': True
            }
            
            buffer = io.BytesIO()
            
            # 3. Al pasarle 'numero_base' (12 dígitos), la librería calcula y agrega el 13vo dígito verificador.
            codigo = CODE_CLASS(numero_base, writer=ImageWriter())
            codigo.write(buffer, options=opciones)
            
            imagen_b64 = base64.b64encode(buffer.getvalue()).decode('utf-8')
            
            descripcion = bien.descripcion_bien or bien.descripcion_item or "Sin descripción"
            
            etiquetas.append({
                # Usamos get_fullcode() para devolver el número completo de 13 dígitos
                "numero": codigo.get_fullcode(), 
                "descripcion": descripcion,
                "imagen": f"data:image/png;base64,{imagen_b64}"
            })
            
        return etiquetas

    def get_tipos_informatica(self):
        # Buscamos todos los tipos que no estén vacíos
        tipos = self.db.query(models.Patrimonio.tipo).filter(
            models.Patrimonio.tipo.isnot(None),
            models.Patrimonio.tipo != ''
        ).distinct().all()
        
        # Filtramos los principales para devolver solo "Los otros"
        principales = ['pc', 'notebook', 'scanner', 'impresora']
        tipos_otros = [t[0] for t in tipos if t[0] and t[0].lower() not in principales]
        
        return tipos_otros

    def get_anios_disponibles(self):
        # Buscamos todos los años que no estén vacíos
        anios = self.db.query(models.Patrimonio.anio).filter(
            models.Patrimonio.anio.isnot(None),
            models.Patrimonio.anio != ''
        ).distinct().order_by(desc(models.Patrimonio.anio)).all()
        
        # Extraemos el primer elemento de la tupla devuelta por SQLAlchemy
        return [a[0] for a in anios if a[0]]

    def generar_pdf_salida(self, data: schemas.NotaSalidaRequest):
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.lib.utils import simpleSplit
        
        # 1. Obtener datos del equipo (automático o manual)
        tipo = data.tipo_manual or "Equipo"
        marca = data.marca_manual or ""
        serie = data.serie_manual or "S/N"
        inventario_str = ""

        if data.numero_inventario:
            # Buscamos si existe en la BD
            bien = self.db.query(models.Patrimonio).filter(models.Patrimonio.numero_inventario == data.numero_inventario).first()
            if bien:
                tipo = bien.tipo or bien.descripcion_bien or "Equipo"
                marca = bien.marca or "Sin marca"
                serie = bien.serie or "S/N"
                inventario_str = f" Nro Inventario: {bien.numero_inventario}"
            else:
                inventario_str = f" Nro Inventario: {data.numero_inventario}"

        # 2. Configurar el Canvas del PDF
        buffer = io.BytesIO()
        c = canvas.Canvas(buffer, pagesize=A4)
        width, height = A4

        # 3. Formatear la fecha en español
        meses = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
        hoy = datetime.now()
        fecha_str = f"Buenos Aires, {hoy.day} de {meses[hoy.month - 1]} de {hoy.year}"

        # --- DIBUJAR EL DOCUMENTO ---
        
        # Cabecera Institucional
        c.setFont("Helvetica-Bold", 14)
        c.drawCentredString(width / 2.0, height - 80, "GOBIERNO DE LA CIUDAD DE BUENOS AIRES")
        c.setFont("Helvetica-Bold", 12)
        c.drawCentredString(width / 2.0, height - 100, "DIRECCIÓN GENERAL TESORERÍA")

        # Fecha
        c.setFont("Helvetica", 11)
        c.drawRightString(width - 70, height - 140, fecha_str)

        # Cuerpo del Texto
        c.setFont("Helvetica", 12)
        
        # Ajustamos el texto exacto según tu modelo
        texto_nota = (
            f"Se hace entrega de {data.cantidad} ({data.cantidad}) {tipo} marca {marca} "
            f"Nro de serie: {serie}{inventario_str} a {data.nombre_apellido} "
            f"CUIT/CUIL {data.cuil_cuit} por el motivo: {data.motivo}."
        )

        # Usamos simpleSplit para que el texto baje de línea automáticamente si es muy largo
        lineas = simpleSplit(texto_nota, "Helvetica", 12, width - 140)
        
        y_text = height - 200
        for linea in lineas:
            c.drawString(70, y_text, linea)
            y_text -= 20

        # Firmas (Líneas y Textos en la parte inferior)
        y_firmas = y_text - 150
        
        # Firma Receptor (Izquierda)
        c.line(90, y_firmas, 240, y_firmas)
        c.setFont("Helvetica-Bold", 10)
        c.drawCentredString(165, y_firmas - 15, "Firma y Aclaración del Receptor")
        
        # Firma Tesorería (Derecha)
        c.line(350, y_firmas, 500, y_firmas)
        c.drawCentredString(425, y_firmas - 15, "Firma Autorizada - Tesorería")

        # Guardar el PDF en memoria
        c.showPage()
        c.save()
        buffer.seek(0)
        return buffer

    def asignar_equipo(self, numero_inventario: str, usuario_id: int, puesto: str, usuario_admin: str, ip: str):
        # 1. Buscamos el equipo
        bien = self.get_by_numero(numero_inventario)
        
        # 2. Buscamos al usuario (si enviaron un ID)
        from backend.modulos.usuarios.models import User
        usuario_asignado = self.db.query(User).filter(User.id == usuario_id).first() if usuario_id else None

        # Guardamos valores anteriores para el historial
        val_ant_usuario = bien.usuario_destino or "Sin asignar"
        val_ant_puesto = bien.puesto or "Sin puesto"

        # 3. Aplicamos la nueva lógica: Custodia + Ubicación
        bien.usuario_id = usuario_id
        if usuario_asignado:
            bien.usuario_destino = f"{usuario_asignado.nombre} {usuario_asignado.apellido} (CUIL: {usuario_asignado.cuil})"
        else:
            bien.usuario_destino = None # Se desasignó
            
        bien.puesto = puesto
        bien.usuario_modificacion = usuario_admin

        # 4. Registramos en el historial
        obs = f"Asignado a: {bien.usuario_destino or 'Nadie'} | Ubicación: {puesto or 'Depósito'}"
        self._registrar_historial(
            numero_inventario=bien.numero_inventario,
            usuario=usuario_admin,
            accion="ASIGNACION",
            campo_modificado="usuario_id/puesto",
            valor_anterior=f"{val_ant_usuario} en {val_ant_puesto}",
            valor_nuevo=obs,
            ip=ip,
            observaciones="Asignación manual desde módulo Informática"
        )

        self.db.commit()
        self.db.refresh(bien)
        return bien

    async def importar_excel_informatica(self, file: UploadFile, usuario: str, ip: str):
        import pandas as pd
        import io
        contenido = await file.read()
        try:
            xls = pd.ExcelFile(io.BytesIO(contenido))
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"No se pudo leer el Excel. Verifica que no esté corrupto. Error: {str(e)}")

        altas, modificaciones, omitidos, total_registros = 0, 0, 0, 0
        bienes_a_crear, historiales = [], []

        def get_val(fila_actual, nombres_columna):
            for col in nombres_columna:
                if col in df.columns:
                    val = fila_actual[col]
                    if pd.isna(val): return None
                    val_str = str(val).strip()
                    if val_str == '' or val_str.lower() == 'nan': return None
                    return val_str
            return None

        for sheet_name in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet_name)
            df.columns = df.columns.astype(str).str.strip()
            df = df.where(pd.notnull(df), None)
            total_registros += len(df)

            for index, row in df.iterrows():
                nro_inv = get_val(row, ['Nº Inventario', 'N  Inventario', 'Nro Inventario', 'numero_inventario', 'Inventario'])
                if not nro_inv:
                    omitidos += 1
                    continue
                
                nro_inv = str(nro_inv).replace(".0", "").strip()

                datos_fila = {
                    "nombre_de_equipo": get_val(row, ['Nombre de Equipo', 'nombre_de_equipo']),
                    "puesto": get_val(row, ['Puesto', 'puesto']),
                    "tipo": get_val(row, ['Tipo', 'tipo']),
                    "marca": get_val(row, ['Marca', 'marca']),
                    "modelo": get_val(row, ['Modelo', 'modelo']),
                    "procesador": get_val(row, ['Procesador', 'procesador']),
                    "motherboard": get_val(row, ['Motherboard', 'motherboard']),
                    "memoria": get_val(row, ['Memoria RAM', 'memoria']),
                    "disco": get_val(row, ['Disco', 'disco']),
                    "serie": get_val(row, ['Serie', 'serie']),
                    "monitor": get_val(row, ['Monitor', 'monitor']),
                    "serie_monitor": get_val(row, ['Serie Monitor', 'serie_monitor']),
                    "usuario": get_val(row, ['Usuario', 'usuario']),
                    "usuario_destino": get_val(row, ['Usuario/Destino', 'usuario_destino']),
                    "observaciones": get_val(row, ['Observaciones', 'observaciones']),
                }
                
                estado_excel = get_val(row, ['Estado', 'estado'])
                datos_fila = {k: v for k, v in datos_fila.items() if v is not None}

                bien_existente = self.db.query(models.Patrimonio).filter(models.Patrimonio.numero_inventario == nro_inv).first()

                if bien_existente:
                    cambios = False
                    if estado_excel and estado_excel.lower() != "autorizado":
                        datos_fila["estado"] = estado_excel
                    
                    for campo, valor_nuevo in datos_fila.items():
                        if getattr(bien_existente, campo) != valor_nuevo:
                            setattr(bien_existente, campo, valor_nuevo)
                            cambios = True
                    
                    if cambios:
                        bien_existente.usuario_modificacion = usuario
                        modificaciones += 1
                        historiales.append(models.PatrimonioHistorial(
                            numero_inventario=nro_inv, usuario=usuario, accion="IMPORTACION_MODIFICACION",
                            ip=ip, observaciones=f"Actualización técnica vía Excel Informática (Hoja {sheet_name})"
                        ))
                    else:
                        omitidos += 1
                else:
                    if estado_excel: datos_fila["estado"] = estado_excel
                    nuevo_bien = models.Patrimonio(
                        numero_inventario=nro_inv, usuario_modificacion=usuario, **datos_fila
                    )
                    bienes_a_crear.append(nuevo_bien)
                    altas += 1
                    historiales.append(models.PatrimonioHistorial(
                        numero_inventario=nro_inv, usuario=usuario, accion="IMPORTACION_ALTA",
                        ip=ip, observaciones=f"Alta técnica vía Excel Informática (Hoja {sheet_name})"
                    ))

        try:
            if bienes_a_crear: self.db.add_all(bienes_a_crear)
            if historiales: self.db.add_all(historiales)
            
            registro_importacion = models.PatrimonioImportacion(
                usuario=usuario, archivo=file.filename, cantidad_registros=total_registros,
                cantidad_altas=altas, cantidad_modificaciones=modificaciones, cantidad_omitidos=omitidos
            )
            self.db.add(registro_importacion)
            self.db.commit()

            return {
                "mensaje": "Importación técnica finalizada",
                "resumen": {
                    "total_filas_leidas": total_registros,
                    "nuevos_creados": altas,
                    "bienes_actualizados": modificaciones,
                    "omitidos_o_sin_cambios": omitidos
                }
            }
        except Exception as e:
            self.db.rollback()
            raise HTTPException(status_code=500, detail=f"Error BD. Detalle: {str(e)}")