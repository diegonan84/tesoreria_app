"""
Mesa de Ayuda (Tickets)
-----------------------
- Cualquier usuario autenticado puede crear tickets y ver/comentar los suyos.
- El personal con rol "Soporte" o "Administrador" ve todos, los asigna,
  cambia estado/prioridad y puede dejar notas internas.
"""
import io
from datetime import datetime, timedelta
from typing import Optional

import pandas as pd
from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import Response
from sqlalchemy import or_, func
from sqlalchemy.orm import Session, joinedload

from backend.database import get_db
from backend.modulos.auth.dependencies import verificar_usuario_autenticado
from backend.modulos.auditoria.service import registrar_evento
from backend.modulos.usuarios.models import User
from backend.modulos.patrimonio.models import Patrimonio

from . import schemas
from .models import Ticket, TicketComentario

router = APIRouter(prefix="/api/tickets", tags=["Mesa de Ayuda"])

ROLES_SOPORTE = {"Soporte", "Administrador"}
ESTADOS_FINALES = {"Resuelto", "Cerrado"}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def es_soporte(usuario: dict) -> bool:
    return bool(ROLES_SOPORTE.intersection(usuario.get("roles", []) or []))


def _requiere_soporte(usuario: dict):
    if not es_soporte(usuario):
        raise HTTPException(status_code=403, detail="Operación permitida solo para Soporte.")


def _auditar(db: Session, usuario: dict, accion: str, detalle: str, request: Request):
    registrar_evento(
        db, usuario_id=usuario.get("id"), accion=accion, detalle=detalle,
        ip_address=request.client.host if request.client else None,
        pc_nombre=request.headers.get("X-PC-Nombre", "Desconocido"),
        pc_usuario=request.headers.get("X-PC-Usuario", "Desconocido"),
    )


def _mini(u: Optional[User]):
    if not u:
        return None
    return schemas.UsuarioMini(
        id=u.id,
        nombre=f"{u.nombre} {u.apellido}".strip(),
        sector=u.sector.nombre if u.sector else None,
    )


def _serializar(t: Ticket, ver_internos: bool) -> schemas.TicketResponse:
    visibles = [c for c in t.comentarios if ver_internos or not c.es_interno]
    return schemas.TicketResponse(
        id=t.id, numero=t.numero, titulo=t.titulo, descripcion=t.descripcion,
        categoria=t.categoria, prioridad=t.prioridad, estado=t.estado,
        numero_inventario=t.numero_inventario,
        fecha_creacion=t.fecha_creacion, fecha_actualizacion=t.fecha_actualizacion,
        fecha_cierre=t.fecha_cierre,
        solicitante=_mini(t.solicitante), asignado=_mini(t.asignado),
        cantidad_comentarios=len([c for c in visibles if c.tipo == "comentario"]),
    )


def _obtener_ticket(db: Session, ticket_id: int, usuario: dict) -> Ticket:
    t = (
        db.query(Ticket)
        .options(joinedload(Ticket.comentarios), joinedload(Ticket.solicitante), joinedload(Ticket.asignado))
        .filter(Ticket.id == ticket_id)
        .first()
    )
    if not t:
        raise HTTPException(status_code=404, detail="Ticket no encontrado")
    if not es_soporte(usuario) and t.solicitante_id != usuario.get("id"):
        # 404 en vez de 403 para no revelar que el ticket existe
        raise HTTPException(status_code=404, detail="Ticket no encontrado")
    return t


def _log_sistema(t: Ticket, usuario_id: int, texto: str):
    t.comentarios.append(TicketComentario(usuario_id=usuario_id, texto=texto, tipo="sistema"))


# ---------------------------------------------------------------------------
# Catálogos y datos auxiliares
# ---------------------------------------------------------------------------
@router.get("/catalogos")
def catalogos(usuario=Depends(verificar_usuario_autenticado)):
    return {
        "categorias": schemas.CATEGORIAS,
        "prioridades": schemas.PRIORIDADES,
        "estados": schemas.ESTADOS,
        "es_soporte": es_soporte(usuario),
    }


@router.get("/agentes")
def listar_agentes(db: Session = Depends(get_db), usuario=Depends(verificar_usuario_autenticado)):
    """Usuarios que pueden recibir tickets (rol Soporte o Administrador)."""
    _requiere_soporte(usuario)
    candidatos = (
        db.query(User)
        .filter(User.activo == True, or_(User.roles.ilike("%Soporte%"), User.roles.ilike("%Administrador%")))
        .order_by(User.apellido, User.nombre)
        .all()
    )
    agentes = []
    for u in candidatos:
        roles = {r.strip() for r in (u.roles or "").split(",")}
        if roles & ROLES_SOPORTE:
            agentes.append(_mini(u))
    return agentes


@router.get("/mis-equipos")
def mis_equipos(db: Session = Depends(get_db), usuario=Depends(verificar_usuario_autenticado)):
    """Equipos de Patrimonio asignados al usuario, para elegir en el alta del ticket."""
    bienes = (
        db.query(Patrimonio)
        .filter(Patrimonio.usuario_id == usuario.get("id"), Patrimonio.activo == True)
        .order_by(Patrimonio.numero_inventario)
        .all()
    )
    return [
        {
            "numero_inventario": b.numero_inventario,
            "descripcion": " ".join(filter(None, [b.tipo or b.descripcion_bien, b.marca, b.modelo])) or "Bien",
            "nombre_de_equipo": b.nombre_de_equipo,
        }
        for b in bienes
    ]


@router.get("/resumen")
def resumen(db: Session = Depends(get_db), usuario=Depends(verificar_usuario_autenticado)):
    """Contadores por estado (para las tarjetas superiores)."""
    q = db.query(Ticket.estado, func.count(Ticket.id))
    if not es_soporte(usuario):
        q = q.filter(Ticket.solicitante_id == usuario.get("id"))
    conteo = {estado: 0 for estado in schemas.ESTADOS}
    for estado, cant in q.group_by(Ticket.estado).all():
        conteo[estado] = cant

    sin_asignar = 0
    mios_asignados = 0
    if es_soporte(usuario):
        abiertos = Ticket.estado.notin_(ESTADOS_FINALES)
        sin_asignar = db.query(Ticket).filter(abiertos, Ticket.asignado_id.is_(None)).count()
        mios_asignados = db.query(Ticket).filter(abiertos, Ticket.asignado_id == usuario.get("id")).count()
    return {"por_estado": conteo, "sin_asignar": sin_asignar, "asignados_a_mi": mios_asignados}


# ---------------------------------------------------------------------------
# Listado / alta / detalle
# ---------------------------------------------------------------------------
def _consulta_listado(db: Session, usuario: dict, alcance: str, estado: Optional[str],
                      prioridad: Optional[str], categoria: Optional[str], q: Optional[str]):
    consulta = db.query(Ticket).options(
        joinedload(Ticket.solicitante).joinedload(User.sector),
        joinedload(Ticket.asignado),
        joinedload(Ticket.comentarios),
    )
    uid = usuario.get("id")

    if not es_soporte(usuario) or alcance == "mios":
        consulta = consulta.filter(Ticket.solicitante_id == uid)
    elif alcance == "asignados":
        consulta = consulta.filter(Ticket.asignado_id == uid)
    elif alcance == "sin_asignar":
        consulta = consulta.filter(Ticket.asignado_id.is_(None))

    if estado == "activos":
        consulta = consulta.filter(Ticket.estado.notin_(ESTADOS_FINALES))
    elif estado:
        consulta = consulta.filter(Ticket.estado == estado)
    if prioridad:
        consulta = consulta.filter(Ticket.prioridad == prioridad)
    if categoria:
        consulta = consulta.filter(Ticket.categoria == categoria)
    if q:
        patron = f"%{q.strip()}%"
        consulta = consulta.filter(or_(
            Ticket.numero.ilike(patron), Ticket.titulo.ilike(patron),
            Ticket.descripcion.ilike(patron), Ticket.numero_inventario.ilike(patron),
        ))
    return consulta.order_by(Ticket.fecha_creacion.desc(), Ticket.id.desc())


@router.get("", response_model=list[schemas.TicketResponse])
def listar_tickets(
    alcance: str = "todos",            # todos | mios | asignados | sin_asignar (solo Soporte)
    estado: Optional[str] = None,      # un estado puntual o "activos"
    prioridad: Optional[str] = None,
    categoria: Optional[str] = None,
    q: Optional[str] = None,
    db: Session = Depends(get_db),
    usuario=Depends(verificar_usuario_autenticado),
):
    tickets = _consulta_listado(db, usuario, alcance, estado, prioridad, categoria, q).limit(500).all()
    ver_internos = es_soporte(usuario)
    return [_serializar(t, ver_internos) for t in tickets]


@router.post("", response_model=schemas.TicketResponse)
def crear_ticket(
    datos: schemas.TicketCreate,
    request: Request,
    db: Session = Depends(get_db),
    usuario=Depends(verificar_usuario_autenticado),
):
    titulo = datos.titulo.strip()
    descripcion = datos.descripcion.strip()
    if len(titulo) < 3 or len(descripcion) < 3:
        raise HTTPException(status_code=400, detail="Completá el título y la descripción.")

    t = Ticket(
        titulo=titulo,
        descripcion=descripcion,
        categoria=datos.categoria,
        prioridad=datos.prioridad,
        numero_inventario=(datos.numero_inventario or "").strip() or None,
        solicitante_id=usuario.get("id"),
        estado="Abierto",
    )
    db.add(t)
    db.flush()  # obtenemos el id para armar el número sin carreras
    t.numero = f"TK-{t.id:05d}"
    _log_sistema(t, usuario.get("id"), "Ticket creado.")
    db.commit()
    db.refresh(t)

    _auditar(db, usuario, "TICKET_CREAR", f"Ticket {t.numero} creado: {t.titulo}", request)
    return _serializar(t, es_soporte(usuario))


@router.get("/exportar")
def exportar_tickets(
    alcance: str = "todos",
    estado: Optional[str] = None,
    prioridad: Optional[str] = None,
    categoria: Optional[str] = None,
    q: Optional[str] = None,
    db: Session = Depends(get_db),
    usuario=Depends(verificar_usuario_autenticado),
):
    _requiere_soporte(usuario)
    tickets = _consulta_listado(db, usuario, alcance, estado, prioridad, categoria, q).all()

    def local(dt):
        # Las fechas se guardan en UTC; Argentina es UTC-3
        return (dt - timedelta(hours=3)).strftime("%d/%m/%Y %H:%M") if dt else ""

    filas = [{
        "Nº": t.numero,
        "Fecha": local(t.fecha_creacion),
        "Título": t.titulo,
        "Categoría": t.categoria,
        "Prioridad": t.prioridad,
        "Estado": t.estado,
        "Solicitante": f"{t.solicitante.nombre} {t.solicitante.apellido}" if t.solicitante else "",
        "Sector": t.solicitante.sector.nombre if t.solicitante and t.solicitante.sector else "",
        "Asignado a": f"{t.asignado.nombre} {t.asignado.apellido}" if t.asignado else "",
        "Nº Inventario": t.numero_inventario or "",
        "Cierre": local(t.fecha_cierre),
        "Descripción": t.descripcion,
    } for t in tickets]

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame(filas or [{"Nº": "Sin tickets"}]).to_excel(writer, index=False, sheet_name="Tickets")
        ws = writer.sheets["Tickets"]
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        for celda in ws[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill(start_color="1A3644", end_color="1A3644", fill_type="solid")
        for idx in range(1, ws.max_column + 1):
            largo = max((len(str(c.value)) for c in ws[get_column_letter(idx)] if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(idx)].width = min(largo + 2, 50)

    return Response(
        buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Tickets_Mesa_de_Ayuda.xlsx"},
    )


@router.get("/{ticket_id}", response_model=schemas.TicketDetalle)
def detalle_ticket(ticket_id: int, db: Session = Depends(get_db), usuario=Depends(verificar_usuario_autenticado)):
    t = _obtener_ticket(db, ticket_id, usuario)
    ver_internos = es_soporte(usuario)
    base = _serializar(t, ver_internos).model_dump()

    comentarios = [
        schemas.ComentarioResponse(
            id=c.id, texto=c.texto, tipo=c.tipo, es_interno=bool(c.es_interno),
            fecha=c.fecha, usuario=_mini(c.usuario),
        )
        for c in t.comentarios if ver_internos or not c.es_interno
    ]

    equipo = None
    if t.numero_inventario:
        b = db.query(Patrimonio).filter(Patrimonio.numero_inventario == t.numero_inventario).first()
        if b:
            equipo = {
                "numero_inventario": b.numero_inventario,
                "tipo": b.tipo or b.descripcion_bien,
                "marca": b.marca, "modelo": b.modelo, "serie": b.serie,
                "nombre_de_equipo": b.nombre_de_equipo, "puesto": b.puesto,
            }

    return schemas.TicketDetalle(**base, comentarios=comentarios, equipo=equipo)


# ---------------------------------------------------------------------------
# Gestión
# ---------------------------------------------------------------------------
@router.put("/{ticket_id}", response_model=schemas.TicketResponse)
def actualizar_ticket(
    ticket_id: int,
    datos: schemas.TicketUpdate,
    request: Request,
    db: Session = Depends(get_db),
    usuario=Depends(verificar_usuario_autenticado),
):
    t = _obtener_ticket(db, ticket_id, usuario)
    uid = usuario.get("id")
    soporte = es_soporte(usuario)
    cambios = []

    if not soporte:
        # El solicitante solo puede cerrar su ticket o reabrirlo si estaba resuelto/cerrado
        if datos.prioridad or datos.categoria or datos.asignado_id or datos.desasignar:
            raise HTTPException(status_code=403, detail="Solo Soporte puede modificar esos datos.")
        if datos.estado not in (None, "Cerrado", "Abierto"):
            raise HTTPException(status_code=403, detail="Solo podés cerrar o reabrir tu ticket.")
        if datos.estado == "Abierto" and t.estado not in ESTADOS_FINALES:
            raise HTTPException(status_code=400, detail="El ticket ya está abierto.")

    if datos.estado and datos.estado != t.estado:
        cambios.append(f"Estado: {t.estado} → {datos.estado}")
        t.estado = datos.estado
        t.fecha_cierre = datetime.utcnow() if datos.estado in ESTADOS_FINALES else None

    if datos.prioridad and datos.prioridad != t.prioridad:
        cambios.append(f"Prioridad: {t.prioridad} → {datos.prioridad}")
        t.prioridad = datos.prioridad

    if datos.categoria and datos.categoria != t.categoria:
        cambios.append(f"Categoría: {t.categoria} → {datos.categoria}")
        t.categoria = datos.categoria

    if datos.desasignar and t.asignado_id:
        cambios.append("Se quitó la asignación")
        t.asignado_id = None
    elif datos.asignado_id and datos.asignado_id != t.asignado_id:
        agente = db.query(User).filter(User.id == datos.asignado_id).first()
        roles_agente = {r.strip() for r in ((agente.roles if agente else "") or "").split(",")}
        if not agente or not (roles_agente & ROLES_SOPORTE):
            raise HTTPException(status_code=400, detail="El usuario elegido no pertenece a Soporte.")
        cambios.append(f"Asignado a {agente.nombre} {agente.apellido}")
        t.asignado_id = agente.id
        # Si se toma un ticket recién abierto, pasa automáticamente a "En curso"
        if t.estado == "Abierto" and not datos.estado:
            cambios.append("Estado: Abierto → En curso")
            t.estado = "En curso"

    if not cambios:
        return _serializar(t, soporte)

    _log_sistema(t, uid, " · ".join(cambios))
    t.fecha_actualizacion = datetime.utcnow()
    db.commit()
    db.refresh(t)

    _auditar(db, usuario, "TICKET_ACTUALIZAR", f"{t.numero}: " + "; ".join(cambios), request)
    return _serializar(t, soporte)


@router.post("/{ticket_id}/tomar", response_model=schemas.TicketResponse)
def tomar_ticket(ticket_id: int, request: Request, db: Session = Depends(get_db),
                 usuario=Depends(verificar_usuario_autenticado)):
    """Atajo para que el agente se asigne el ticket a sí mismo."""
    _requiere_soporte(usuario)
    return actualizar_ticket(ticket_id, schemas.TicketUpdate(asignado_id=usuario.get("id")), request, db, usuario)


@router.post("/{ticket_id}/comentarios", response_model=schemas.ComentarioResponse)
def comentar_ticket(
    ticket_id: int,
    datos: schemas.ComentarioCreate,
    request: Request,
    db: Session = Depends(get_db),
    usuario=Depends(verificar_usuario_autenticado),
):
    t = _obtener_ticket(db, ticket_id, usuario)
    texto = datos.texto.strip()
    if not texto:
        raise HTTPException(status_code=400, detail="El comentario está vacío.")
    soporte = es_soporte(usuario)
    interno = bool(datos.es_interno) and soporte

    c = TicketComentario(ticket_id=t.id, usuario_id=usuario.get("id"), texto=texto,
                         tipo="comentario", es_interno=interno)
    db.add(c)

    # Si el solicitante responde a un ticket "En espera", vuelve a "En curso"
    if not soporte and t.estado == "En espera":
        t.estado = "En curso"
        _log_sistema(t, usuario.get("id"), "El solicitante respondió · Estado: En espera → En curso")
    t.fecha_actualizacion = datetime.utcnow()

    db.commit()
    db.refresh(c)

    _auditar(db, usuario, "TICKET_COMENTAR", f"Comentario{' interno' if interno else ''} en {t.numero}", request)
    autor = db.query(User).filter(User.id == usuario.get("id")).first()
    return schemas.ComentarioResponse(id=c.id, texto=c.texto, tipo=c.tipo, es_interno=c.es_interno,
                                      fecha=c.fecha, usuario=_mini(autor))
